"""Relatórios (RF09, RF10), pesquisa por aluno (RF11) e respostas por atividade (RF12)."""
import re
from dataclasses import dataclass, field
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy.orm import joinedload

from .. import db
from ..models import Atividade, CursoTurma, Resposta
from .cadastros import chave_nome
from .presenca import normalizar_matricula
from .tempo import agora, formatar_data, formatar_data_hora
from .validacao import sem_acentos, slug

MAX_ALUNOS_PESQUISA = 50
_NEGRITO = Font(bold=True)
_CARACTERES_PROIBIDOS_ABA = re.compile(r"[\\/?*\[\]:]")


def _ordem_atividade(atividade):
    return (atividade.data, atividade.hora_inicio, chave_nome(atividade.titulo))


def rotulo_atividade(atividade):
    """Ex.: "20/10/2026 – Título"."""
    return f"{formatar_data(atividade.data)} – {atividade.titulo}"


# --- Planilhas: utilidades ---------------------------------------------------

def _anexar(aba, valores):
    """Acrescenta uma linha. Texto que começa com "=" fica como texto, nunca fórmula:
    nome e descrição vêm do formulário público."""
    aba.append(valores)
    for celula in aba[aba.max_row]:
        if isinstance(celula.value, str) and celula.value.startswith("="):
            celula.data_type = "s"


def _cabecalho(aba, titulos, larguras):
    _anexar(aba, titulos)
    for celula, largura in zip(aba[1], larguras):
        celula.font = _NEGRITO
        celula.alignment = Alignment(vertical="top", wrap_text=True)
        aba.column_dimensions[celula.column_letter].width = largura


def _salvar(livro):
    saida = BytesIO()
    livro.save(saida)
    saida.seek(0)
    return saida


def nome_aba(nome, usados):
    """Nome válido e único de aba do Excel: até 31 caracteres, sem / \\ ? * [ ] :."""
    base = _CARACTERES_PROIBIDOS_ABA.sub("-", nome).strip(" '") or "Turma"
    candidato, n = base[:31], 2
    while candidato.casefold() in usados:
        sufixo = f" ({n})"
        candidato, n = base[:31 - len(sufixo)] + sufixo, n + 1
    usados.add(candidato.casefold())
    return candidato


# --- RF10: consolidado por curso/turma ---------------------------------------

@dataclass
class AlunoEvento:
    """Uma linha do consolidado: um aluno de uma turma em um evento."""
    matricula: str
    evento: str
    nome: str = ""
    atividades: list = field(default_factory=list)  # Atividade, por data


@dataclass
class TurmaConsolidada:
    nome: str
    linhas: list  # AlunoEvento, por nome do aluno e nome do evento

    @property
    def quantidade_alunos(self):
        return len({linha.matricula for linha in self.linhas})

    @property
    def quantidade_presencas(self):
        return sum(len(linha.atividades) for linha in self.linhas)


def consolidar_por_turma(tipo_evento_id=None):
    """Agrupa as respostas por curso/turma e, dentro dela, por matrícula e evento.

    Um aluno que informou turmas diferentes aparece em cada uma delas, contando
    só as atividades daquela turma. O nome vem da resposta mais recente do aluno
    na turma.
    """
    consulta = (
        db.select(Resposta, CursoTurma.nome)
        .join(Resposta.curso_turma)
        .join(Resposta.atividade)
        .options(joinedload(Resposta.atividade).joinedload(Atividade.tipo_evento))
    )
    if tipo_evento_id:
        consulta = consulta.where(Atividade.tipo_evento_id == tipo_evento_id)

    turmas = {}  # nome da turma -> (linhas por (matrícula, evento), nome mais recente)
    for resposta, nome_turma in db.session.execute(consulta):
        linhas, nomes = turmas.setdefault(nome_turma, ({}, {}))
        atividade = resposta.atividade
        chave = (resposta.matricula, atividade.tipo_evento_id)
        linha = linhas.setdefault(
            chave, AlunoEvento(resposta.matricula, atividade.tipo_evento.nome))
        linha.atividades.append(atividade)
        envio, _ = nomes.get(resposta.matricula, (None, None))
        if envio is None or resposta.enviado_em >= envio:
            nomes[resposta.matricula] = (resposta.enviado_em, resposta.nome)

    resultado = []
    for nome_turma, (linhas, nomes) in sorted(turmas.items(),
                                             key=lambda item: chave_nome(item[0])):
        for linha in linhas.values():
            linha.nome = nomes[linha.matricula][1]
            linha.atividades.sort(key=_ordem_atividade)
        resultado.append(TurmaConsolidada(nome_turma, sorted(
            linhas.values(),
            key=lambda l: (chave_nome(l.nome), l.matricula, chave_nome(l.evento)))))
    return resultado


def texto_atividades(atividades):
    """Ex.: "20/10/2026 – Abertura; 21/10/2026 – Robótica"."""
    return "; ".join(rotulo_atividade(atividade) for atividade in atividades)


def exportar_consolidado(turmas, tipo_evento=None):
    livro = Workbook()
    resumo = livro.active
    resumo.title = "Resumo"
    usados = {"resumo"}
    filtro = tipo_evento.nome if tipo_evento else "Todos os eventos"
    _anexar(resumo, ["Consolidado de presença por curso/turma"])
    resumo["A1"].font = Font(bold=True, size=13)
    _anexar(resumo, ["Tipo de evento", filtro])
    _anexar(resumo, ["Gerado em", formatar_data_hora(agora())])
    _anexar(resumo, [])
    _anexar(resumo, ["Curso/turma", "Aba", "Alunos", "Presenças"])
    for celula in resumo[5]:
        celula.font = _NEGRITO

    for turma in turmas:
        aba = livro.create_sheet(nome_aba(turma.nome, usados))
        _anexar(resumo, [turma.nome, aba.title, turma.quantidade_alunos,
                         turma.quantidade_presencas])

        _cabecalho(aba, ["Nome", "Matrícula", "Evento", "Atividades",
                         "Quantidade de atividades"], [36, 16, 24, 80, 14])
        for linha in turma.linhas:
            _anexar(aba, [linha.nome, linha.matricula, linha.evento,
                          texto_atividades(linha.atividades), len(linha.atividades)])
            aba.cell(aba.max_row, 4).alignment = Alignment(wrap_text=True, vertical="top")
            aba.cell(aba.max_row, 5).alignment = Alignment(horizontal="center",
                                                           vertical="top")
        aba.freeze_panes = "A2"
        aba.auto_filter.ref = aba.dimensions

    for coluna, largura in zip("ABCD", [36, 24, 10, 12]):
        resumo.column_dimensions[coluna].width = largura
    return _salvar(livro)


def nome_arquivo_consolidado(tipo_evento=None):
    if tipo_evento is None:
        return "consolidado-por-turma-todos-os-eventos.xlsx"
    return f"consolidado-por-turma-{slug(tipo_evento.nome, 'evento')}.xlsx"


# --- RF09 e RF12: respostas de uma atividade ---------------------------------

def respostas_da_atividade(atividade):
    return db.session.execute(
        db.select(Resposta).filter_by(atividade_id=atividade.id)
        .options(joinedload(Resposta.curso_turma))
        .order_by(Resposta.enviado_em, Resposta.id)
    ).scalars().all()


def exportar_respostas(atividade, respostas):
    livro = Workbook()
    aba = livro.active
    aba.title = "Presença"
    _cabecalho(aba, ["Nome", "Matrícula", "Curso/turma", "E-mail", "Descrição",
                     "Enviado em"], [36, 16, 24, 32, 60, 17])
    for r in respostas:
        _anexar(aba, [r.nome, r.matricula, r.curso_turma.nome, r.email, r.descricao,
                      r.enviado_em])
        aba.cell(aba.max_row, 6).number_format = "dd/mm/yyyy hh:mm"
        aba.cell(aba.max_row, 5).alignment = Alignment(wrap_text=True, vertical="top")
    aba.freeze_panes = "A2"
    return _salvar(livro)


def nome_arquivo_respostas(atividade):
    return f"presenca-{slug(atividade.titulo, 'atividade')[:60].strip('-')}-{atividade.id}.xlsx"


# --- RF11: pesquisa por aluno ------------------------------------------------

@dataclass
class AlunoEncontrado:
    matricula: str
    nome: str = ""
    ultimo_envio: object = None
    cursos: set = field(default_factory=set)
    presencas: list = field(default_factory=list)  # Resposta, por data da atividade


def _chave_busca(texto):
    return chave_nome(sem_acentos(texto))


def pesquisar_alunos(termo):
    """Busca por matrícula (se o termo tiver números) ou por nome, sem acentos.

    Retorna (alunos, mais_resultados): no máximo MAX_ALUNOS_PESQUISA alunos.
    """
    termo = (termo or "").strip()
    consulta = db.select(Resposta).options(
        joinedload(Resposta.atividade).joinedload(Atividade.tipo_evento),
        joinedload(Resposta.curso_turma),
    )
    if any(c.isdigit() for c in termo):
        matricula = normalizar_matricula(termo)
        if len(matricula) < 2:
            return [], False
        respostas = db.session.execute(
            consulta.where(Resposta.matricula.contains(matricula, autoescape=True))
        ).scalars().all()
    else:
        procurado = _chave_busca(termo)
        if len(procurado) < 2:
            return [], False
        respostas = [r for r in db.session.execute(consulta).scalars()
                     if procurado in _chave_busca(r.nome)]

    alunos = {}
    for resposta in respostas:
        aluno = alunos.setdefault(resposta.matricula, AlunoEncontrado(resposta.matricula))
        if aluno.ultimo_envio is None or resposta.enviado_em >= aluno.ultimo_envio:
            aluno.nome, aluno.ultimo_envio = resposta.nome, resposta.enviado_em
        aluno.cursos.add(resposta.curso_turma.nome)
        aluno.presencas.append(resposta)
    for aluno in alunos.values():
        aluno.presencas.sort(key=lambda r: _ordem_atividade(r.atividade))

    ordenados = sorted(alunos.values(), key=lambda a: (chave_nome(a.nome), a.matricula))
    return ordenados[:MAX_ALUNOS_PESQUISA], len(ordenados) > MAX_ALUNOS_PESQUISA
