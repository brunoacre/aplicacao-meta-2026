"""Importação de atividades por planilha XLSX (RF06, formato na seção 6 do SPEC).

Fluxo: o arquivo enviado é guardado temporariamente, lido e validado linha a
linha para a pré-visualização; na confirmação, é lido e validado de novo (o
banco pode ter mudado) e só as linhas válidas são gravadas.
"""
import re
import time as relogio
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from .. import db
from ..models import MODALIDADES, Atividade, TipoEvento
from .cadastros import (
    chave_nome, chaves_atividades, listar_tipos_evento, tipo_evento_por_nome,
)
from .tempo import fechamento_padrao
from .validacao import sem_acentos

CABECALHOS = [
    "tipo_evento", "atividade", "modalidade", "data",
    "hora_inicio", "local", "envolvidos",
]
# Planilhas feitas antes da retirada da hora de término: a coluna é ignorada.
CABECALHOS_ANTIGOS = CABECALHOS[:5] + ["hora_fim"] + CABECALHOS[5:]
MAX_LINHAS = 1000
# Arquivos de pré-visualização não confirmados são apagados depois deste tempo.
VALIDADE_ARQUIVO = timedelta(days=1)

# Limites iguais aos do cadastro manual (modelos e AtividadeForm).
_TAMANHO_MAXIMO = {"tipo_evento": 120, "atividade": 300, "local": 120}
_HORA = re.compile(r"^(\d{1,2})[:hH](\d{2})$")
_DATA = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})$")


class ErroPlanilha(Exception):
    """Problema que impede a leitura da planilha inteira (arquivo ou cabeçalho)."""


@dataclass
class LinhaImportacao:
    numero: int
    valores: dict  # texto como veio na planilha, para exibir na pré-visualização
    dados: dict = field(default_factory=dict)  # valores já convertidos
    erros: list = field(default_factory=list)
    novo_tipo: bool = False  # o tipo de evento será criado na importação

    @property
    def valida(self):
        return not self.erros


@dataclass
class ResultadoLeitura:
    linhas: list

    @property
    def validas(self):
        return [linha for linha in self.linhas if linha.valida]

    @property
    def com_erro(self):
        return [linha for linha in self.linhas if not linha.valida]


# --- Modelo em branco --------------------------------------------------------

def gerar_modelo():
    """Planilha em branco com os cabeçalhos e uma aba de instruções."""
    livro = Workbook()
    aba = livro.active
    aba.title = "Atividades"
    aba.append(CABECALHOS)
    for celula, largura in zip(aba[1], [24, 50, 16, 12, 12, 20, 40]):
        celula.font = Font(bold=True)
        aba.column_dimensions[celula.column_letter].width = largura
    aba.freeze_panes = "A2"

    instrucoes = livro.create_sheet("Instruções")
    for linha in [
        ["Preencha uma atividade por linha na aba Atividades, sem alterar os cabeçalhos."],
        ["tipo_evento: nome do evento, ex.: 35ª META 2026. Se não existir, será criado."],
        ["atividade: título da atividade."],
        ["modalidade: palestra, apresentacao, minicurso ou outra."],
        ["data: DD/MM/AAAA."],
        ["hora_inicio: HH:MM. O formulário de presença abre à 00h00 da data."],
        ["local e envolvidos podem ficar em branco; os demais campos são obrigatórios."],
        ["Atividades com o mesmo título no mesmo evento não são importadas."],
    ]:
        instrucoes.append(linha)
    instrucoes.column_dimensions["A"].width = 90

    saida = BytesIO()
    livro.save(saida)
    saida.seek(0)
    return saida


# --- Conversão de células ----------------------------------------------------

def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


# Aceita o código (apresentacao) e o rótulo (Apresentação de trabalho).
_MODALIDADES_ACEITAS = {
    chave_nome(sem_acentos(texto)): codigo
    for codigo, rotulo in MODALIDADES.items()
    for texto in (codigo, rotulo)
}


def converter_modalidade(valor):
    return _MODALIDADES_ACEITAS.get(chave_nome(sem_acentos(_texto(valor))))


def converter_data(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    encontrado = _DATA.match(_texto(valor))
    if encontrado:
        dia, mes, ano = map(int, encontrado.groups())
        try:
            return date(ano, mes, dia)
        except ValueError:
            return None
    return None


def converter_hora(valor):
    if isinstance(valor, datetime):
        valor = valor.time()
    if isinstance(valor, timedelta) and timedelta(0) <= valor < timedelta(days=1):
        valor = (datetime.min + valor).time()
    if isinstance(valor, time):
        return valor.replace(second=0, microsecond=0)
    encontrado = _HORA.match(_texto(valor))
    if encontrado:
        hora, minuto = map(int, encontrado.groups())
        if hora < 24 and minuto < 60:
            return time(hora, minuto)
    return None


# --- Leitura e validação -----------------------------------------------------

def _abrir(arquivo):
    try:
        livro = load_workbook(arquivo, read_only=True, data_only=True)
    except Exception as erro:  # arquivo corrompido, outro formato etc.
        raise ErroPlanilha(
            "Não foi possível ler o arquivo. Envie uma planilha no formato .xlsx.") from erro
    return livro


def ler_planilha(arquivo):
    """Lê a primeira aba e valida cada linha. Levanta ErroPlanilha se não der para ler."""
    livro = _abrir(arquivo)
    try:
        linhas = list(livro.worksheets[0].iter_rows(values_only=True))
    finally:
        livro.close()

    if not linhas:
        raise ErroPlanilha("A planilha está vazia.")
    cabecalho = [_texto(v).casefold() for v in linhas[0]]
    while cabecalho and not cabecalho[-1]:
        cabecalho.pop()
    if cabecalho == CABECALHOS_ANTIGOS:
        colunas = [CABECALHOS_ANTIGOS.index(nome) for nome in CABECALHOS]
    elif cabecalho == CABECALHOS:
        colunas = list(range(len(CABECALHOS)))
    else:
        raise ErroPlanilha(
            "Os cabeçalhos da primeira linha devem ser exatamente: "
            + ", ".join(CABECALHOS) + ". Baixe o modelo para conferir.")

    corpo = [
        (numero, [linha[i] if i < len(linha) else None for i in colunas])
        for numero, linha in enumerate(linhas[1:], start=2)
        if any(_texto(v) for v in linha)
    ]
    if not corpo:
        raise ErroPlanilha("A planilha não tem nenhuma atividade preenchida.")
    if len(corpo) > MAX_LINHAS:
        raise ErroPlanilha(f"A planilha tem mais de {MAX_LINHAS} atividades. "
                           "Divida-a em arquivos menores.")

    return validar_linhas(corpo)


def _validar_linha(numero, celulas):
    celulas = list(celulas) + [None] * (len(CABECALHOS) - len(celulas))
    brutos = dict(zip(CABECALHOS, celulas))
    valores = {nome: _texto(valor) for nome, valor in brutos.items()}
    linha = LinhaImportacao(numero=numero, valores=valores)
    erros = linha.erros

    for nome in ("tipo_evento", "atividade"):
        if not valores[nome]:
            erros.append(f"Informe o campo {nome}.")
    for nome, limite in _TAMANHO_MAXIMO.items():
        if len(valores[nome]) > limite:
            erros.append(f"O campo {nome} passa de {limite} caracteres.")

    modalidade = converter_modalidade(brutos["modalidade"])
    if not valores["modalidade"]:
        erros.append("Informe o campo modalidade.")
    elif modalidade is None:
        erros.append("Modalidade inválida: use palestra, apresentacao, minicurso ou outra.")

    dia = converter_data(brutos["data"])
    if not valores["data"]:
        erros.append("Informe o campo data.")
    elif dia is None:
        erros.append("Data inválida: use DD/MM/AAAA.")

    hora_inicio = converter_hora(brutos["hora_inicio"])
    if not valores["hora_inicio"]:
        erros.append("Informe o campo hora_inicio.")
    elif hora_inicio is None:
        erros.append("Hora inválida em hora_inicio: use HH:MM.")

    linha.dados = {
        "tipo_evento": " ".join(valores["tipo_evento"].split()),
        "titulo": valores["atividade"],
        "modalidade": modalidade,
        "data": dia,
        "hora_inicio": hora_inicio,
        "local": valores["local"],
        "envolvidos": valores["envolvidos"],
    }
    return linha


def validar_linhas(corpo):
    """Valida as linhas (número, células) e aplica a regra de duplicidade."""
    existentes = chaves_atividades()
    tipos_existentes = {chave_nome(t.nome) for t in listar_tipos_evento()}
    vistas = {}  # chave -> número da primeira linha válida com ela
    linhas = []
    for numero, celulas in corpo:
        linha = _validar_linha(numero, celulas)
        linhas.append(linha)
        if not linha.valida:
            continue

        chave = (chave_nome(linha.dados["tipo_evento"]), chave_nome(linha.dados["titulo"]))
        if chave in existentes:
            linha.erros.append("Já existe uma atividade com este título neste evento.")
        elif chave in vistas:
            linha.erros.append(f"Atividade repetida na planilha (igual à linha {vistas[chave]}).")
        else:
            vistas[chave] = numero
            linha.novo_tipo = chave[0] not in tipos_existentes
    return ResultadoLeitura(linhas)


# --- Gravação ----------------------------------------------------------------

def importar(resultado):
    """Grava as linhas válidas numa única transação. Retorna (atividades, tipos criados)."""
    tipos = {}
    criados = []
    quantidade = 0
    for linha in resultado.validas:
        dados = linha.dados
        chave = chave_nome(dados["tipo_evento"])
        if chave not in tipos:
            tipo = tipo_evento_por_nome(dados["tipo_evento"])
            if tipo is None:
                tipo = TipoEvento(nome=dados["tipo_evento"])
                db.session.add(tipo)
                criados.append(tipo.nome)
            tipos[chave] = tipo
        db.session.add(Atividade(
            tipo_evento=tipos[chave],
            titulo=dados["titulo"],
            modalidade=dados["modalidade"],
            data=dados["data"],
            hora_inicio=dados["hora_inicio"],
            local=dados["local"],
            envolvidos=dados["envolvidos"],
            fecha_em=fechamento_padrao(dados["data"]),
        ))
        quantidade += 1
    db.session.commit()
    return quantidade, criados


# --- Arquivo temporário da pré-visualização ----------------------------------

def guardar_arquivo(arquivo, pasta):
    """Salva o arquivo enviado e retorna o identificador usado na confirmação."""
    pasta = Path(pasta)
    pasta.mkdir(parents=True, exist_ok=True)
    _remover_antigos(pasta)
    identificador = uuid.uuid4().hex
    arquivo.save(pasta / f"{identificador}.xlsx")
    return identificador


def caminho_arquivo(pasta, identificador):
    """Caminho do arquivo guardado, ou None se o identificador for inválido ou expirou."""
    if not re.fullmatch(r"[0-9a-f]{32}", identificador or ""):
        return None
    caminho = Path(pasta) / f"{identificador}.xlsx"
    return caminho if caminho.is_file() else None


def _remover_antigos(pasta):
    limite = relogio.time() - VALIDADE_ARQUIVO.total_seconds()
    for caminho in pasta.glob("*.xlsx"):
        if caminho.stat().st_mtime < limite:
            caminho.unlink(missing_ok=True)
