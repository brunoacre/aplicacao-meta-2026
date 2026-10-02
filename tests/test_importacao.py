from datetime import date, datetime, time
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app import db
from app.models import Atividade, TipoEvento
from app.servicos import importacao
from app.servicos.importacao import CABECALHOS, ErroPlanilha, ler_planilha
from app.servicos.tempo import fechamento_padrao


def _planilha(*linhas, cabecalhos=CABECALHOS):
    livro = Workbook()
    aba = livro.active
    aba.append(cabecalhos)
    for linha in linhas:
        aba.append(list(linha))
    saida = BytesIO()
    livro.save(saida)
    saida.seek(0)
    return saida


def _linha(**extra):
    dados = {
        "tipo_evento": "35ª META 2026", "atividade": "Palestra de abertura",
        "modalidade": "palestra", "data": "20/10/2026",
        "hora_inicio": "14:00", "local": "Auditório", "envolvidos": "Fulano, Beltrana",
    }
    dados.update(extra)
    return [dados[c] for c in CABECALHOS]


def _ler(app, *linhas, **kwargs):
    with app.app_context():
        return ler_planilha(_planilha(*linhas, **kwargs))


def _criar_atividade(app, nome_evento, titulo):
    with app.app_context():
        tipo = TipoEvento(nome=nome_evento)
        dia = date(2026, 10, 20)
        db.session.add(Atividade(
            tipo_evento=tipo, titulo=titulo, modalidade="palestra", data=dia,
            hora_inicio=time(14), fecha_em=fechamento_padrao(dia)))
        db.session.commit()


def _enviar(cliente, *linhas):
    return cliente.post("/admin/importar", data={
        "arquivo": (_planilha(*linhas), "atividades.xlsx")},
        content_type="multipart/form-data")


def _identificador(html):
    return html.split('name="identificador" type="hidden" value="')[1].split('"')[0]


# --- leitura e validação -----------------------------------------------------

def test_linha_valida_e_convertida(app):
    resultado = _ler(app, _linha())
    linha, = resultado.validas
    assert linha.numero == 2
    assert linha.novo_tipo is True
    assert linha.dados == {
        "tipo_evento": "35ª META 2026", "titulo": "Palestra de abertura",
        "modalidade": "palestra", "data": date(2026, 10, 20),
        "hora_inicio": time(14), "local": "Auditório", "envolvidos": "Fulano, Beltrana",
    }


def test_celulas_de_data_e_hora_do_excel(app):
    resultado = _ler(app, _linha(data=datetime(2026, 10, 20),
                                 hora_inicio=datetime(1899, 12, 30, 14, 20)))
    linha, = resultado.validas
    assert linha.dados["data"] == date(2026, 10, 20)
    assert linha.dados["hora_inicio"] == time(14, 20)


@pytest.mark.parametrize("texto, codigo", [
    ("PALESTRA", "palestra"),
    ("Apresentação", "apresentacao"),
    (" apresentacao ", "apresentacao"),
    ("Apresentação de trabalho", "apresentacao"),
    ("MiniCurso", "minicurso"),
    ("outra", "outra"),
])
def test_modalidade_sem_diferenciar_maiusculas_e_acentos(app, texto, codigo):
    linha, = _ler(app, _linha(modalidade=texto)).validas
    assert linha.dados["modalidade"] == codigo


def test_local_e_envolvidos_podem_ficar_em_branco(app):
    linha, = _ler(app, _linha(local=None, envolvidos=None)).validas
    assert linha.dados["local"] == ""
    assert linha.dados["envolvidos"] == ""


@pytest.mark.parametrize("campo, valor, mensagem", [
    ("tipo_evento", None, "Informe o campo tipo_evento."),
    ("atividade", "  ", "Informe o campo atividade."),
    ("modalidade", None, "Informe o campo modalidade."),
    ("modalidade", "oficina", "Modalidade inválida"),
    ("data", None, "Informe o campo data."),
    ("data", "2026-10-20", "Data inválida"),
    ("data", "31/02/2026", "Data inválida"),
    ("hora_inicio", None, "Informe o campo hora_inicio."),
    ("hora_inicio", "25:00", "Hora inválida em hora_inicio"),
    ("hora_inicio", "duas horas", "Hora inválida em hora_inicio"),
    ("atividade", "x" * 301, "O campo atividade passa de 300 caracteres."),
    ("local", "x" * 121, "O campo local passa de 120 caracteres."),
])
def test_erros_por_linha(app, campo, valor, mensagem):
    resultado = _ler(app, _linha(), _linha(**{"atividade": "Outra", campo: valor}))
    assert len(resultado.validas) == 1
    linha, = resultado.com_erro
    assert linha.numero == 3
    assert any(mensagem in erro for erro in linha.erros)


def test_linhas_vazias_sao_ignoradas(app):
    resultado = _ler(app, _linha(), [None] * 7, ["", " "], _linha(atividade="Outra"))
    assert [linha.numero for linha in resultado.linhas] == [2, 5]


def test_cabecalho_diferente(app):
    cabecalhos = CABECALHOS[:3] + ["dia"] + CABECALHOS[4:]
    with pytest.raises(ErroPlanilha, match="cabeçalhos"):
        _ler(app, _linha(), cabecalhos=cabecalhos)


def test_planilha_antiga_com_hora_fim_e_aceita(app):
    # A3: a coluna hora_fim de planilhas antigas é ignorada.
    antiga = _linha()
    antiga.insert(5, "15:00")
    linha, = _ler(app, antiga, cabecalhos=importacao.CABECALHOS_ANTIGOS).validas
    assert linha.dados["hora_inicio"] == time(14)
    assert linha.dados["local"] == "Auditório"
    assert linha.dados["envolvidos"] == "Fulano, Beltrana"


def test_modelo_nao_tem_hora_fim(app):
    aba = load_workbook(importacao.gerar_modelo()).active
    assert [c.value for c in aba[1]] == CABECALHOS
    assert "hora_fim" not in CABECALHOS


def test_cabecalho_com_maiusculas_e_espacos_e_aceito(app):
    cabecalhos = [f" {c.upper()} " for c in CABECALHOS]
    assert len(_ler(app, _linha(), cabecalhos=cabecalhos).validas) == 1


def test_planilha_sem_atividades(app):
    with pytest.raises(ErroPlanilha, match="nenhuma atividade"):
        _ler(app)


def test_arquivo_que_nao_e_xlsx(app):
    with app.app_context(), pytest.raises(ErroPlanilha, match="formato .xlsx"):
        ler_planilha(BytesIO(b"isto nao e uma planilha"))


# --- duplicidade -------------------------------------------------------------

def test_duplicada_com_o_banco_pelo_nome_do_evento(app):
    _criar_atividade(app, "35ª META 2026", "Palestra de abertura")
    resultado = _ler(app, _linha(tipo_evento=" 35ª meta  2026", atividade="PALESTRA de abertura"))
    linha, = resultado.com_erro
    assert linha.erros == ["Já existe uma atividade com este título neste evento."]


def test_mesmo_titulo_em_outro_evento_e_permitido(app):
    _criar_atividade(app, "Semana C&T 2026", "Palestra de abertura")
    assert len(_ler(app, _linha()).validas) == 1


def test_duplicada_dentro_da_planilha(app):
    resultado = _ler(app, _linha(), _linha(atividade=" palestra DE abertura"))
    linha, = resultado.com_erro
    assert linha.numero == 3
    assert linha.erros == ["Atividade repetida na planilha (igual à linha 2)."]


# --- gravação ----------------------------------------------------------------

def test_importar_cria_tipo_e_grava_so_validas(app):
    with app.app_context():
        resultado = ler_planilha(_planilha(
            _linha(), _linha(atividade="Minicurso", modalidade="minicurso"),
            _linha(atividade="Com erro", data="ontem")))
        assert importacao.importar(resultado) == (2, ["35ª META 2026"])

        tipo = db.session.execute(db.select(TipoEvento)).scalar_one()
        assert (tipo.nome, tipo.ativo) == ("35ª META 2026", True)
        atividades = db.session.execute(db.select(Atividade)).scalars().all()
        assert sorted(a.titulo for a in atividades) == ["Minicurso", "Palestra de abertura"]
        assert all(a.tipo_evento_id == tipo.id for a in atividades)
        assert all(a.fecha_em == datetime(2026, 10, 20, 23, 59) for a in atividades)
        assert len({a.token for a in atividades}) == 2


def test_importar_usa_tipo_existente_pelo_nome(app):
    with app.app_context():
        existente = TipoEvento(nome="35ª META 2026")
        outro = TipoEvento(nome="34ª META 2025")
        db.session.add_all([outro, existente])
        db.session.commit()
        resultado = ler_planilha(_planilha(_linha(tipo_evento=" 35ª meta  2026 ")))
        assert resultado.validas[0].novo_tipo is False
        assert importacao.importar(resultado) == (1, [])
        atividade = db.session.execute(db.select(Atividade)).scalar_one()
        assert atividade.tipo_evento_id == existente.id
        assert db.session.query(TipoEvento).count() == 2


# --- rotas -------------------------------------------------------------------

@pytest.mark.parametrize("metodo, url", [
    ("get", "/admin/importar"),
    ("post", "/admin/importar"),
    ("get", "/admin/importar/modelo"),
    ("post", "/admin/importar/confirmar"),
])
def test_rotas_exigem_login(client, metodo, url):
    resposta = getattr(client, metodo)(url)
    assert resposta.status_code == 302
    assert "/admin/login" in resposta.headers["Location"]


def test_baixar_modelo(cliente_logado):
    resposta = cliente_logado.get("/admin/importar/modelo")
    assert resposta.status_code == 200
    assert "modelo_atividades.xlsx" in resposta.headers["Content-Disposition"]
    aba = load_workbook(BytesIO(resposta.data)).worksheets[0]
    assert [c.value for c in aba[1]] == CABECALHOS
    assert aba.max_row == 1


def test_modelo_e_aceito_pela_leitura(app):
    with app.app_context(), pytest.raises(ErroPlanilha, match="nenhuma atividade"):
        ler_planilha(importacao.gerar_modelo())


def test_previa_e_confirmacao(app, cliente_logado):
    _criar_atividade(app, "35ª META 2026", "Já cadastrada")
    resposta = _enviar(cliente_logado, _linha(), _linha(atividade="Já cadastrada"))
    html = resposta.get_data(as_text=True)
    assert "Pré-visualização da importação" in html
    assert "Importar 1 atividade válida" in html
    assert "Já existe uma atividade com este título neste evento." in html
    with app.app_context():
        assert db.session.query(Atividade).count() == 1  # nada gravado ainda

    identificador = _identificador(html)
    resposta = cliente_logado.post("/admin/importar/confirmar",
                                   data={"identificador": identificador}, follow_redirects=True)
    html = resposta.get_data(as_text=True)
    assert "1 atividade importada." in html
    assert "1 linha(s) com erro não foram importadas." in html
    with app.app_context():
        assert db.session.query(Atividade).count() == 2
    assert importacao.caminho_arquivo(
        app.config["PASTA_IMPORTACOES"], identificador) is None


def test_confirmacao_revalida_a_planilha(app, cliente_logado):
    html = _enviar(cliente_logado, _linha()).get_data(as_text=True)
    # Alguém cadastra a mesma atividade entre a prévia e a confirmação.
    _criar_atividade(app, "35ª META 2026", "Palestra de abertura")
    resposta = cliente_logado.post("/admin/importar/confirmar",
                                   data={"identificador": _identificador(html)},
                                   follow_redirects=True)
    assert "0 atividades importadas." in resposta.get_data(as_text=True)
    with app.app_context():
        assert db.session.query(Atividade).count() == 1


@pytest.mark.parametrize("identificador", ["", "../../etc/passwd", "0" * 32])
def test_confirmacao_com_arquivo_invalido_ou_expirado(cliente_logado, identificador):
    resposta = cliente_logado.post("/admin/importar/confirmar",
                                   data={"identificador": identificador}, follow_redirects=True)
    assert "A pré-visualização expirou." in resposta.get_data(as_text=True)


def test_envio_com_cabecalho_errado_mostra_erro(cliente_logado):
    resposta = cliente_logado.post("/admin/importar", data={
        "arquivo": (_planilha(_linha(), cabecalhos=["a", "b"]), "atividades.xlsx")},
        content_type="multipart/form-data")
    assert "Os cabeçalhos da primeira linha devem ser exatamente" in resposta.get_data(
        as_text=True)


def test_envio_de_arquivo_com_outra_extensao(cliente_logado):
    resposta = cliente_logado.post("/admin/importar", data={
        "arquivo": (BytesIO(b"a;b"), "atividades.csv")}, content_type="multipart/form-data")
    assert "Envie uma planilha no formato .xlsx." in resposta.get_data(as_text=True)


def test_envio_de_arquivo_grande_demais(app, cliente_logado):
    app.config["MAX_CONTENT_LENGTH"] = 1024
    resposta = cliente_logado.post("/admin/importar", data={
        "arquivo": (BytesIO(b"x" * 4096), "atividades.xlsx")},
        content_type="multipart/form-data", follow_redirects=True)
    assert "O arquivo passa do tamanho máximo de 2 MB." in resposta.get_data(as_text=True)
