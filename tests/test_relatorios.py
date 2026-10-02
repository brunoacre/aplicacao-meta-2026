from datetime import date, datetime, time
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app import db
from app.models import Atividade, CursoTurma, Resposta, TipoEvento
from app.servicos import relatorios
from app.servicos.tempo import fechamento_padrao

DESCRICAO = "Gostei muito da atividade. " * 5


@pytest.fixture
def base(app):
    """Dois eventos, três atividades e duas turmas; devolve os IDs."""
    with app.app_context():
        meta = TipoEvento(nome="35ª META 2026")
        semana = TipoEvento(nome="Semana C&T 2026")
        cursos = [CursoTurma(nome="Informática 1/A"), CursoTurma(nome="Mecatrônica 2B")]

        def atividade(tipo, titulo, dia, hora):
            return Atividade(tipo_evento=tipo, titulo=titulo, modalidade="palestra",
                             data=dia, hora_inicio=time(hora), fecha_em=fechamento_padrao(dia))
        atividades = [
            atividade(meta, "Abertura", date(2026, 10, 20), 9),
            atividade(meta, "Robótica", date(2026, 10, 21), 14),
            atividade(semana, "Astronomia", date(2026, 11, 5), 19),
        ]
        db.session.add_all([meta, semana, *cursos, *atividades])
        db.session.commit()
        return {"meta": meta.id, "semana": semana.id,
                "info": cursos[0].id, "meca": cursos[1].id,
                "abertura": atividades[0].id, "robotica": atividades[1].id,
                "astronomia": atividades[2].id}


def _responder(app, atividade_id, curso_id, matricula, nome, email="aluno@email.com",
               enviado_em=datetime(2026, 10, 20, 10), descricao=DESCRICAO):
    with app.app_context():
        db.session.add(Resposta(atividade_id=atividade_id, curso_turma_id=curso_id,
                                matricula=matricula, nome=nome, email=email,
                                descricao=descricao, enviado_em=enviado_em, ip="1.2.3.4"))
        db.session.commit()


def _planilha(resposta):
    return load_workbook(BytesIO(resposta.data))


def _linhas(aba):
    """Valores de cada linha, sem as células vazias do fim."""
    linhas = []
    for linha in aba.iter_rows(values_only=True):
        linha = list(linha)
        while linha and linha[-1] is None:
            linha.pop()
        linhas.append(linha)
    return linhas


# --- acesso ------------------------------------------------------------------

@pytest.mark.parametrize("url", [
    "/admin/relatorios",
    "/admin/relatorios/consolidado.xlsx",
    "/admin/atividades/1/respostas",
    "/admin/atividades/1/respostas.xlsx",
])
def test_rotas_exigem_login(client, url):
    resposta = client.get(url)
    assert resposta.status_code == 302
    assert "/admin/login" in resposta.headers["Location"]


# --- RF10: consolidado por turma ---------------------------------------------

def test_consolidado_uma_aba_por_turma_com_lista_de_alunos(app, cliente_logado, base):
    _responder(app, base["robotica"], base["info"], "2026001", "Bruna Silva")
    _responder(app, base["abertura"], base["info"], "2026001", "Bruna Silva")
    _responder(app, base["abertura"], base["info"], "2026002", "Ana Souza")
    _responder(app, base["astronomia"], base["meca"], "2026003", "Carlos Lima")

    resposta = cliente_logado.get("/admin/relatorios/consolidado.xlsx")
    assert resposta.status_code == 200
    assert ("consolidado-por-turma-todos-os-eventos.xlsx"
            in resposta.headers["Content-Disposition"])
    livro = _planilha(resposta)
    assert livro.sheetnames == ["Resumo", "Informática 1-A", "Mecatrônica 2B"]

    # A2: nome, matrícula, evento, atividades concatenadas (por data) e quantidade.
    linhas = _linhas(livro["Informática 1-A"])
    assert linhas == [
        ["Nome", "Matrícula", "Evento", "Atividades", "Quantidade de atividades"],
        ["Ana Souza", "2026002", "35ª META 2026", "20/10/2026 – Abertura", 1],
        ["Bruna Silva", "2026001", "35ª META 2026",
         "20/10/2026 – Abertura; 21/10/2026 – Robótica", 2],
    ]
    assert _linhas(livro["Mecatrônica 2B"])[1] == [
        "Carlos Lima", "2026003", "Semana C&T 2026", "05/11/2026 – Astronomia", 1]

    resumo = _linhas(livro["Resumo"])
    assert ["Tipo de evento", "Todos os eventos"] in resumo
    assert ["Informática 1/A", "Informática 1-A", 2, 3] in resumo
    assert ["Mecatrônica 2B", "Mecatrônica 2B", 1, 1] in resumo


def test_consolidado_sem_filtro_tem_uma_linha_por_evento(app, base):
    # A quantidade conta só as atividades de cada evento.
    _responder(app, base["abertura"], base["info"], "2026001", "Bruna")
    _responder(app, base["robotica"], base["info"], "2026001", "Bruna")
    _responder(app, base["astronomia"], base["info"], "2026001", "Bruna")
    with app.app_context():
        turma, = relatorios.consolidar_por_turma()
        assert [(l.evento, len(l.atividades)) for l in turma.linhas] == [
            ("35ª META 2026", 2), ("Semana C&T 2026", 1)]
        assert (turma.quantidade_alunos, turma.quantidade_presencas) == (1, 3)


def test_consolidado_filtrado_por_evento(app, cliente_logado, base):
    _responder(app, base["abertura"], base["info"], "2026001", "Bruna")
    _responder(app, base["astronomia"], base["info"], "2026001", "Bruna")
    _responder(app, base["astronomia"], base["meca"], "2026003", "Carlos")

    resposta = cliente_logado.get(f"/admin/relatorios/consolidado.xlsx?tipo_evento={base['meta']}")
    assert "consolidado-por-turma-35a-meta-2026.xlsx" in resposta.headers["Content-Disposition"]
    livro = _planilha(resposta)
    assert livro.sheetnames == ["Resumo", "Informática 1-A"]
    assert _linhas(livro["Informática 1-A"])[1:] == [
        ["Bruna", "2026001", "35ª META 2026", "20/10/2026 – Abertura", 1]]


def test_aluno_com_turmas_diferentes_aparece_em_cada_uma(app, base):
    _responder(app, base["abertura"], base["info"], "2026001", "Bruna")
    _responder(app, base["robotica"], base["meca"], "2026001", "Bruna")
    with app.app_context():
        turmas = relatorios.consolidar_por_turma()
        assert [(t.nome, [[a.id for a in l.atividades] for l in t.linhas])
                for t in turmas] == [
            ("Informática 1/A", [[base["abertura"]]]),
            ("Mecatrônica 2B", [[base["robotica"]]]),
        ]


def test_nome_da_resposta_mais_recente(app, base):
    _responder(app, base["robotica"], base["info"], "2026001", "Bruna S. Lima",
               enviado_em=datetime(2026, 10, 21, 15))
    _responder(app, base["abertura"], base["info"], "2026001", "bruna",
               enviado_em=datetime(2026, 10, 20, 10))
    _responder(app, base["astronomia"], base["info"], "2026001", "bruna",
               enviado_em=datetime(2026, 11, 5, 9))
    with app.app_context():
        linhas = relatorios.consolidar_por_turma()[0].linhas
        # Mesmo nome em todas as linhas do aluno: o da resposta mais recente na turma.
        assert [l.nome for l in linhas] == ["bruna", "bruna"]


def test_planilha_nao_tem_email(app, cliente_logado, base):
    _responder(app, base["abertura"], base["info"], "2026001", "Bruna",
               email="bruna@email.com")
    livro = _planilha(cliente_logado.get("/admin/relatorios/consolidado.xlsx"))
    assert "bruna@email.com" not in str(_linhas(livro["Informática 1-A"]))


def test_consolidado_sem_respostas(cliente_logado, base):
    resposta = cliente_logado.get("/admin/relatorios/consolidado.xlsx", follow_redirects=True)
    assert "Nenhuma presença registrada para gerar o relatório." in resposta.get_data(
        as_text=True)


@pytest.mark.parametrize("nome, usados, esperado", [
    ("Informática 1/A", set(), "Informática 1-A"),
    ("Turma [noite]: a*b?", set(), "Turma -noite-- a-b-"),
    ("x" * 40, set(), "x" * 31),
    ("Resumo", {"resumo"}, "Resumo (2)"),
    ("x" * 40, {"x" * 31}, "x" * 27 + " (2)"),
    ("'''", set(), "Turma"),
])
def test_nome_aba(nome, usados, esperado):
    assert relatorios.nome_aba(nome, usados) == esperado


# --- RF09 e RF12: respostas de uma atividade ---------------------------------

def test_tela_de_respostas_somente_leitura(app, cliente_logado, base):
    _responder(app, base["abertura"], base["info"], "2026001", "Bruna Silva")
    texto = cliente_logado.get(f"/admin/atividades/{base['abertura']}/respostas").get_data(
        as_text=True)
    assert "Bruna Silva" in texto and "2026001" in texto and "Gostei muito" in texto
    assert "Exportar XLSX" in texto
    assert "<form" not in texto.split("</header>")[1]  # nenhuma ação sobre as respostas


def test_tela_de_respostas_vazia_e_inexistente(cliente_logado, base):
    texto = cliente_logado.get(f"/admin/atividades/{base['abertura']}/respostas").get_data(
        as_text=True)
    assert "Nenhuma resposta registrada para esta atividade." in texto
    assert cliente_logado.get("/admin/atividades/999/respostas").status_code == 404


def test_exportar_respostas_da_atividade(app, cliente_logado, base):
    _responder(app, base["abertura"], base["info"], "2026002", "Ana",
               enviado_em=datetime(2026, 10, 20, 10, 30))
    _responder(app, base["abertura"], base["meca"], "2026001", "=HYPERLINK(\"x\")",
               enviado_em=datetime(2026, 10, 20, 10, 5))
    _responder(app, base["robotica"], base["info"], "2026009", "Outra atividade")

    resposta = cliente_logado.get(f"/admin/atividades/{base['abertura']}/respostas.xlsx")
    assert f"presenca-abertura-{base['abertura']}.xlsx" in resposta.headers[
        "Content-Disposition"]
    aba = _planilha(resposta).active
    linhas = _linhas(aba)
    assert linhas[0] == ["Nome", "Matrícula", "Curso/turma", "E-mail", "Descrição",
                         "Enviado em"]
    assert linhas[1] == ['=HYPERLINK("x")', "2026001", "Mecatrônica 2B", "aluno@email.com",
                         DESCRICAO, datetime(2026, 10, 20, 10, 5)]
    assert linhas[2][:2] == ["Ana", "2026002"]
    assert len(linhas) == 3
    assert aba["A2"].data_type == "s"  # texto do aluno nunca vira fórmula


# --- RF11: pesquisa por aluno ------------------------------------------------

def test_pesquisa_por_matricula_com_pontos_e_tracos(app, cliente_logado, base):
    _responder(app, base["abertura"], base["info"], "20260015", "João Pereira")
    _responder(app, base["astronomia"], base["meca"], "20260015", "João Pereira")
    _responder(app, base["abertura"], base["info"], "20269999", "Outro")
    texto = cliente_logado.get("/admin/relatorios?q=2026.001-5").get_data(as_text=True)
    assert "João Pereira" in texto and "Outro" not in texto
    assert "2 atividades" in texto
    assert "Abertura" in texto and "Astronomia" in texto
    assert "Semana C&amp;T 2026" in texto


def test_pesquisa_por_nome_sem_acentos(app, base):
    _responder(app, base["abertura"], base["info"], "1", "João Pereira")
    _responder(app, base["abertura"], base["info"], "2", "Maria Joana")
    _responder(app, base["abertura"], base["info"], "3", "Carlos")
    with app.app_context():
        alunos, mais = relatorios.pesquisar_alunos("  JOA ")
        assert [a.nome for a in alunos] == ["João Pereira", "Maria Joana"]
        assert mais is False


def test_pesquisa_limita_resultados(app, base, monkeypatch):
    monkeypatch.setattr(relatorios, "MAX_ALUNOS_PESQUISA", 2)
    for n in range(3):
        _responder(app, base["abertura"], base["info"], str(n), f"Aluno {n}")
    with app.app_context():
        alunos, mais = relatorios.pesquisar_alunos("aluno")
        assert len(alunos) == 2 and mais is True


@pytest.mark.parametrize("termo", ["", "a", "1-.", "%"])
def test_pesquisa_termo_curto_ou_vazio(app, base, termo):
    _responder(app, base["abertura"], base["info"], "2026001", "Ana")
    with app.app_context():
        assert relatorios.pesquisar_alunos(termo) == ([], False)


def test_pagina_de_relatorios(cliente_logado, base):
    texto = cliente_logado.get("/admin/relatorios?q=ninguem").get_data(as_text=True)
    assert "Exportar consolidado por turma" in texto
    assert "35ª META 2026" in texto
    assert "Nenhum aluno encontrado" in texto
